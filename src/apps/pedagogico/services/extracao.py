"""Extração de texto em memória para arquivos PDF e DOCX.

LGPD / Segurança:
- O arquivo enviado NUNCA é gravado em disco nem persistido no banco.
- Limite estrito de 20 MB.
- Apenas arquivos PDF (.pdf) e DOCX (.docx) são aceitos.
"""

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_TAMANHO_BYTES = 20 * 1024 * 1024  # 20 MB

MIME_TYPES_PERMITIDOS = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}

EXTENSOES_PERMITIDAS = {
    ".pdf": "pdf",
    ".docx": "docx",
}


class ExtracaoError(Exception):
    """Erro de validação ou extração do arquivo enviado."""


def identificar_origem(arquivo) -> str:
    """Identifica se o arquivo é 'pdf' ou 'docx' pelo nome ou content_type.

    Levanta ExtracaoError se o formato não for suportado.
    """
    nome = getattr(arquivo, "name", "") or ""
    extensao = Path(nome).suffix.lower()
    content_type = (getattr(arquivo, "content_type", "") or "").lower().strip()

    if extensao in EXTENSOES_PERMITIDAS:
        origem = EXTENSOES_PERMITIDAS[extensao]
        if (
            content_type
            and content_type not in ("application/octet-stream", "")
            and content_type in MIME_TYPES_PERMITIDOS
            and MIME_TYPES_PERMITIDOS[content_type] != origem
        ):
            raise ExtracaoError("O conteúdo do arquivo não corresponde à extensão informada.")
        return origem

    if content_type in MIME_TYPES_PERMITIDOS:
        return MIME_TYPES_PERMITIDOS[content_type]

    raise ExtracaoError("Formato de arquivo inválido. Envie apenas arquivos PDF ou DOCX.")


def extrair_texto(arquivo) -> str:
    """Extrai o texto de um arquivo PDF ou DOCX inteiramente em memória.

    Valida presença, tamanho (máximo 20 MB) e formato antes de processar.
    Levanta ExtracaoError em qualquer falha de validação ou leitura.
    """
    if arquivo is None:
        raise ExtracaoError("Nenhum arquivo foi enviado.")

    tamanho = getattr(arquivo, "size", None)
    if tamanho is not None and tamanho > MAX_TAMANHO_BYTES:
        raise ExtracaoError("O arquivo excede o limite máximo de 20 MB.")

    origem = identificar_origem(arquivo)

    if hasattr(arquivo, "seek"):
        arquivo.seek(0)
    conteudo_bytes = arquivo.read()

    if not conteudo_bytes:
        raise ExtracaoError("O arquivo enviado está vazio.")

    if len(conteudo_bytes) > MAX_TAMANHO_BYTES:
        raise ExtracaoError("O arquivo excede o limite máximo de 20 MB.")

    stream = BytesIO(conteudo_bytes)

    if origem == "pdf":
        texto = _extrair_pdf(stream)
    else:
        texto = _extrair_docx(stream)

    texto_limpo = texto.strip()
    if not texto_limpo:
        raise ExtracaoError("Não foi possível extrair texto do arquivo enviado.")

    return texto_limpo


def _extrair_pdf(stream: BytesIO) -> str:
    try:
        leitor = PdfReader(stream)
        paginas = [pagina.extract_text() or "" for pagina in leitor.pages]
    except Exception as erro:
        raise ExtracaoError("Não foi possível ler o arquivo PDF enviado.") from erro
    return "\n\n".join(trecho.strip() for trecho in paginas if trecho and trecho.strip())


def _extrair_docx(stream: BytesIO) -> str:
    try:
        documento = Document(stream)
        partes: list[str] = []
        for paragrafo in documento.paragraphs:
            if paragrafo.text and paragrafo.text.strip():
                partes.append(paragrafo.text.strip())
        for tabela in documento.tables:
            for linha in tabela.rows:
                celulas = [c.text.strip() for c in linha.cells if c.text and c.text.strip()]
                if celulas:
                    partes.append(" | ".join(celulas))
    except Exception as erro:
        raise ExtracaoError("Não foi possível ler o arquivo DOCX enviado.") from erro
    return "\n\n".join(partes)
